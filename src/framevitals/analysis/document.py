"""Bounded, non-executing inspection of common document file types.

This is document *diagnostics*, not OCR, document conversion, or automatic
classification. No macros, scripts, attachments, or embedded programs run.
"""

from __future__ import annotations

import hashlib
import re
from collections import Counter
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from framevitals.core.beacons import beacon
from framevitals.quality_results import DriftResult
from framevitals.result import AnalysisResult


_MAX_BYTES = 96 * 1024 * 1024
_MAX_TEXT_BYTES = 8 * 1024 * 1024
_WORDS = re.compile(r"\b[\w'-]+\b", re.UNICODE)
_SPACES = re.compile(r"\s+")


class _HTMLText(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.fragments: list[str] = []
        self.skip = 0
        self.headings = 0
        self.links = 0
        self.tags = Counter()

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        self.tags[tag] += 1
        if tag in {"script", "style"}:
            self.skip += 1
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            self.headings += 1
        if tag == "a":
            self.links += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"} and self.skip:
            self.skip -= 1

    def handle_data(self, data: str) -> None:
        if not self.skip and data.strip():
            self.fragments.append(data)


def _decode_text(path: Path) -> tuple[str, bool]:
    if path.stat().st_size > _MAX_TEXT_BYTES:
        raise ValueError("Text file exceeds 8 MiB parsing limit.")
    raw = path.read_bytes()
    if b"\x00" in raw[:2048]:
        raise ValueError("File contains binary NUL bytes; it is not plain text.")
    try:
        return raw.decode("utf-8-sig"), False
    except UnicodeDecodeError:
        return raw.decode("utf-8", errors="replace"), True


def _extract(path: Path, suffix: str, depth: str) -> tuple[str, dict[str, Any]]:
    text_cap = {
        "quick": 120_000,
        "standard": 500_000,
        "deep": 1_500_000,
        "research": 2_000_000,
    }.get(depth, 500_000)
    meta: dict[str, Any] = {"format": suffix.lstrip("."), "truncated": False}

    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ImportError('PDF support: pip install "framevitals[documents]"') from exc
        reader = PdfReader(path, strict=False)
        if reader.is_encrypted:
            raise ValueError("Encrypted PDFs require a password; automatic decryption is disabled.")
        count = len(reader.pages)
        page_limit = {"quick": 16, "standard": 64, "deep": 200, "research": 400}.get(
            depth, 64
        )
        chunks = []
        empty_pages = 0
        for page in list(reader.pages)[:page_limit]:
            extracted = (page.extract_text() or "")[:80_000]
            if not extracted.strip():
                empty_pages += 1
            chunks.append(extracted)
            if sum(map(len, chunks)) >= text_cap:
                meta["truncated"] = True
                break
        text = "\n".join(chunks)[:text_cap]
        meta.update({
            "pages": count,
            "pages_analyzed": len(chunks),
            "pages_without_text": empty_pages,
            "truncated": meta["truncated"] or len(chunks) < count,
            "encrypted": False,
            "ocr_performed": False,
        })
        return text, meta

    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise ImportError('DOCX support: pip install "framevitals[documents]"') from exc
        doc = Document(path)
        lines = []
        headings = 0
        for paragraph in doc.paragraphs[:15_000]:
            if paragraph.style and paragraph.style.name.lower().startswith("heading"):
                headings += 1
            if paragraph.text.strip():
                lines.append(paragraph.text)
            if sum(map(len, lines)) >= text_cap:
                meta["truncated"] = True
                break
        meta.update({
            "paragraphs": len(doc.paragraphs),
            "headings": headings,
            "tables": len(doc.tables),
            "sections": len(doc.sections),
        })
        return "\n".join(lines)[:text_cap], meta

    if suffix == ".pptx":
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise ImportError('PPTX support: pip install "framevitals[documents]"') from exc
        presentation = Presentation(path)
        lines = []
        slides = len(presentation.slides)
        slide_limit = {"quick": 30, "standard": 100, "deep": 300, "research": 500}.get(
            depth, 100
        )
        slides_seen = 0
        for slide in list(presentation.slides)[:slide_limit]:
            slides_seen += 1
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    lines.append(shape.text)
            if sum(map(len, lines)) >= text_cap:
                meta["truncated"] = True
                break
        meta.update({
            "slides": slides,
            "slides_analyzed": slides_seen,
            "truncated": meta["truncated"] or slides_seen < slides,
        })
        return "\n".join(lines)[:text_cap], meta

    if suffix in {".html", ".htm"}:
        source, bad_encoding = _decode_text(path)
        parser = _HTMLText()
        parser.feed(source)
        meta.update({
            "headings": parser.headings,
            "links": parser.links,
            "tags": dict(parser.tags.most_common(25)),
            "encoding_replacements": bad_encoding,
        })
        full = "\n".join(parser.fragments)
        meta["truncated"] = len(full) > text_cap
        return full[:text_cap], meta

    if suffix == ".xml":
        try:
            from defusedxml.ElementTree import iterparse
        except ImportError as exc:
            raise ImportError('XML support: pip install "framevitals[documents]"') from exc
        fragments: list[str] = []
        elements = 0
        chars = 0
        with path.open("rb") as stream:
            for _, element in iterparse(stream, events=("end",)):
                elements += 1
                content = (element.text or "").strip()
                if content and chars < text_cap:
                    fragments.append(content[:text_cap - chars])
                    chars += len(content)
                element.clear()
                if elements > 300_000 or chars >= text_cap:
                    meta["truncated"] = True
                    break
        meta["xml_elements"] = elements
        return "\n".join(fragments)[:text_cap], meta

    full, bad_encoding = _decode_text(path)
    meta["encoding_replacements"] = bad_encoding
    if suffix in {".md", ".markdown"}:
        meta["headings"] = sum(
            line.lstrip().startswith("#") for line in full.splitlines()
        )
    if suffix in {".py", ".c", ".cpp", ".h", ".hpp", ".js", ".ts", ".tsx"}:
        meta["code_lines"] = len(full.splitlines())
    meta["truncated"] = len(full) > text_cap
    return full[:text_cap], meta


def inspect_document(path: str | Path, *, depth: str | None = None) -> dict[str, Any]:
    source = Path(path).expanduser()
    suffix = source.suffix.lower()
    if not source.is_file():
        raise FileNotFoundError(source)
    size = source.stat().st_size
    if size > _MAX_BYTES:
        raise ValueError("Document exceeds 96 MiB inspection limit.")
    mode = str(depth or "standard").lower()
    content, meta = _extract(source, suffix, mode)
    paragraphs = [_SPACES.sub(" ", value).strip()
                  for value in content.splitlines() if value.strip()]
    counts = Counter(paragraphs)
    duplicates = sum(count - 1 for count in counts.values() if count > 1)
    words = _WORDS.findall(content)
    replacement_count = content.count("\ufffd")
    return {
        **meta,
        "filename": source.name,
        "size_bytes": size,
        "characters": len(content),
        "words": len(words),
        "lines": len(content.splitlines()),
        "unique_words": len(set(word.lower() for word in words)),
        "duplicate_lines": duplicates,
        "replacement_characters": replacement_count,
        "empty_text": not bool(content.strip()),
        "sampled_text_sha256": hashlib.sha256(content.encode("utf-8")).hexdigest(),
        "sampled": bool(meta["truncated"]),
    }


def analyze_document(path: str | Path, *, depth: str | None = None) -> AnalysisResult:
    info = inspect_document(path, depth=depth)
    findings: list[dict[str, Any]] = []
    if info["empty_text"]:
        findings.append(beacon(
            "document.no_extracted_text",
            "No extractable text was found",
            severity="medium",
            summary="This document has no extractable text in the inspected portion.",
            recommendation="If this is a scanned PDF, use a trusted OCR workflow separately.",
            evidence={"format": info["format"], "pages_analyzed": info.get("pages_analyzed")},
        ))
    if info.get("pages_analyzed", 0) >= 4:
        if info["pages_without_text"] / info["pages_analyzed"] >= 0.5:
            findings.append(beacon(
                "document.suspected_scanned_pages",
                "Many PDF pages contain no extractable text",
                severity="medium",
                summary="Image-only pages may require OCR; OCR was not attempted.",
                evidence={
                    "pages_without_text": info["pages_without_text"],
                    "pages_analyzed": info["pages_analyzed"],
                },
            ))
    if info["replacement_characters"]:
        findings.append(beacon(
            "document.encoding_replacements",
            "Invalid text encoding was replaced",
            severity="medium",
            summary=f"{info['replacement_characters']} replacement characters were observed.",
            recommendation="Check the document encoding and extraction pipeline.",
        ))
    if info["truncated"]:
        findings.append(beacon(
            "document.partial_inspection",
            "Document was inspected within a bounded budget",
            severity="info",
            summary="Results cover only the analyzed portion of this document.",
            recommendation="Increase Prism depth for additional coverage.",
        ))
    score = 100.0
    if info["empty_text"]:
        score -= 30.0
    if info["replacement_characters"]:
        score -= 10.0
    if info["truncated"]:
        score -= 5.0
    return AnalysisResult({
        "dataset_id": None,
        "filename": info["filename"],
        "analysis_mode": str(depth or "standard"),
        "source_kind": "document",
        "profile": {
            "shape": {"rows": info["lines"], "columns": 1},
            "structure": info,
        },
        "document": info,
        "health": {"overall_score": score,
                   "label": "healthy" if score >= 90 else "attention",
                   "score_basis": "heuristic", "calibrated": False},
        "ml_readiness": {"score": None, "label": "not_applicable"},
        "findings": findings,
        "artifacts_enabled": False,
        "execution": {
            "method": "bounded_document_text_inspection",
            "resource_bounded": True,
            "sampled": info["sampled"],
            "ocr_performed": False,
        },
    })


def compare_documents(reference: str | Path, current: str | Path) -> DriftResult:
    before = inspect_document(reference, depth="standard")
    after = inspect_document(current, depth="standard")
    word_change = abs(after["words"] - before["words"]) / max(1, before["words"])
    line_change = abs(after["lines"] - before["lines"]) / max(1, before["lines"])
    score = max(min(1.0, word_change), min(1.0, line_change))
    if before["sampled_text_sha256"] != after["sampled_text_sha256"]:
        score = max(score, 0.15)
    if before["format"] != after["format"]:
        score = max(score, 0.8)
    severity = ("severe" if score >= 0.75 else
                "moderate" if score >= 0.4 else
                "minor" if score >= 0.15 else "stable")
    status = ("fail" if severity == "severe" else
              "warn" if severity != "stable" else "pass")
    return DriftResult({
        "available": True,
        "source_kind": "document",
        "gate": {"status": status, "severity": severity},
        "summary": {
            "overall_verdict": severity,
            "change_score": round(score, 6),
            "reference_words": before["words"],
            "current_words": after["words"],
        },
        "document": {
            "reference": before,
            "current": after,
            "text_changed": before["sampled_text_sha256"] != after["sampled_text_sha256"],
            "comparison_scope": "extracted_text_metadata",
        },
        "columns": [],
    })
