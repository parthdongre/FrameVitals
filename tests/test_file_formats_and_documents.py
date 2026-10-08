"""Real CLI and document-format tests.

These tests ensure recognized formats reach the intended protocol instead of
falling through to the CSV/tabular loader.
"""

from __future__ import annotations

import json
import sys

import numpy as np
import pytest

import framevitals as fv
from framevitals.cli import main
from framevitals.file_formats import (
    check_file_format,
    format_catalog,
    prepare_file_source,
    render_formats,
    supported_extensions,
)
from framevitals.terminal_ui import TerminalState, run_diagnostic, view_lines


def test_formats_registry_includes_pdf_and_all_core_sources():
    extensions = set(supported_extensions())
    assert {
        ".pdf", ".docx", ".pptx", ".txt", ".md", ".html", ".xml",
        ".json", ".jsonl", ".yaml", ".toml", ".ipynb", ".npy",
        ".onnx", ".safetensors", ".graphml", ".gml", ".csv", ".xlsx",
    } <= extensions
    assert "PDF document" in render_formats()
    assert any(item["category"] == "Document" for item in format_catalog())


def test_cli_formats_command_runs_without_tty(monkeypatch, capsys):
    monkeypatch.setattr(sys, "argv", ["framevitals", "formats"])
    assert main() == 0
    text = capsys.readouterr().out
    assert ".pdf" in text
    assert ".docx" in text

    monkeypatch.setattr(sys, "argv", ["framevitals", "formats", "--json"])
    assert main() == 0
    payload = json.loads(capsys.readouterr().out)
    assert any(".pdf" in item["extensions"] for item in payload)


def test_tui_formats_page_uses_same_registry():
    state = TerminalState(page=3)
    content = "\n".join(view_lines(state))
    assert ".pdf" in content
    assert ".onnx" in content


def test_unknown_and_unsafe_files_fail_early(tmp_path):
    source = tmp_path / "unknown.xyz"
    source.write_bytes(b"content")
    with pytest.raises(ValueError, match="framevitals formats"):
        check_file_format(source)
    unsafe = tmp_path / "model.pkl"
    unsafe.write_bytes(b"not pickle")
    with pytest.raises(ValueError, match="Unsafe pickle"):
        prepare_file_source(unsafe)


def test_plain_text_markdown_html_and_cpp_prism(tmp_path):
    texts = {
        "draft.txt": "FrameVitals detects structural issues.\nMore text here.\n",
        "guide.md": "# Chapter\n\nThis is an example paragraph.\n",
        "page.html": "<html><h1>About</h1><p>Some text</p><script>SECRET</script></html>",
        "main.cpp": "#include <iostream>\nint main() {return 0;}\n",
    }
    for filename, source in texts.items():
        path = tmp_path / filename
        path.write_text(source, encoding="utf-8")
        report = fv.prism(path, depth="quick")
        assert report.analysis["source_kind"] == "document"
        assert report.analysis["document"]["words"] > 0
        assert "DOCUMENT" in report.analysis.summary_text()
        assert "Document overview" in report.analysis.to_html()
        if filename.endswith(".html"):
            assert report.analysis["document"]["headings"] == 1


def test_document_tide_axiom_and_pulse_work(tmp_path):
    baseline = tmp_path / "baseline.txt"
    current = tmp_path / "current.txt"
    baseline.write_text("A working document with words.\n", encoding="utf-8")
    current.write_text("A working document with words and extra details.\n", encoding="utf-8")

    comparison = fv.tide(baseline, current)
    assert comparison["source_kind"] == "document"
    assert comparison["document"]["text_changed"]

    contract = fv.axiom(baseline)
    assert contract.contract["source_kind"] == "document"

    snapshot = fv.pulse(current)
    assert snapshot["state"]["source_kind"] == "document"
    assert snapshot["state"]["structured"]["words"] > 0

    report = fv.prism(current, reference=baseline, depth="quick")
    assert report.analysis["source_kind"] == "document"
    assert report.change is not None
    assert report.validation is not None


def test_npy_direct_cli_prism_routes_to_tensor(tmp_path, monkeypatch, capsys):
    path = tmp_path / "matrix.npy"
    np.save(path, np.eye(5))
    monkeypatch.setattr(
        sys, "argv", ["framevitals", "prism", str(path), "--depth", "quick"],
    )
    assert main() == 0
    assert "TENSOR" in capsys.readouterr().out


def test_jsonl_yaml_toml_loaders(tmp_path):
    jsonl = tmp_path / "items.jsonl"
    jsonl.write_text('{"x":1}\n{"x":2}\n', encoding="utf-8")
    assert prepare_file_source(jsonl) == [{"x": 1}, {"x": 2}]
    assert fv.prism(jsonl, depth="quick").analysis["source_kind"] == "nested"

    toml = tmp_path / "project.toml"
    toml.write_text("[tool]\nmode = 'quick'\n", encoding="utf-8")
    assert fv.prism(toml, depth="quick").analysis["source_kind"] == "nested"

    pytest.importorskip("yaml")
    yaml_path = tmp_path / "cfg.yaml"
    yaml_path.write_text("threshold: 10\n", encoding="utf-8")
    assert prepare_file_source(yaml_path) == {"threshold": 10}


def test_pdf_textless_page_emits_ocr_warning(tmp_path):
    pypdf = pytest.importorskip("pypdf")
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=612, height=792)
    path = tmp_path / "scan.pdf"
    with path.open("wb") as out:
        writer.write(out)

    result = run_diagnostic("Prism", path, depth="quick")
    assert result.result.analysis["source_kind"] == "document"
    document = result.result.analysis["document"]
    assert document["pages"] == 1
    assert document["ocr_performed"] is False
    assert document["empty_text"] is True
    assert any(item["code"] == "document.no_extracted_text" for item in result.findings)


def test_docx_document_structure(tmp_path):
    docx = pytest.importorskip("docx")
    doc = docx.Document()
    doc.add_heading("Report", level=1)
    doc.add_paragraph("FrameVitals inspects paragraphs and document metadata.")
    doc.add_table(rows=2, cols=2)
    path = tmp_path / "report.docx"
    doc.save(path)

    result = fv.prism(path, depth="quick")
    assert result.analysis["document"]["tables"] == 1
    assert result.analysis["document"]["headings"] >= 1


def test_pptx_slide_structure(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text = "FrameVitals presentation content"
    path = tmp_path / "slides.pptx"
    presentation.save(path)

    result = fv.prism(path, depth="quick")
    assert result.analysis["document"]["slides"] == 1
    assert result.analysis["document"]["words"] >= 2
