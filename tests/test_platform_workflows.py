"""Cross-interface and security regressions for the shared platform boundary."""

import io
import json
import subprocess
import sys
from pathlib import Path
from zipfile import ZipFile

import numpy as np
import pytest

from framevitals import capabilities, prism
from framevitals.input_safety import validate_input
from framevitals.result import AnalysisResult
from framevitals.result_contract import build_envelope


def test_result_json_is_strict_and_preserves_missing_values():
    result = AnalysisResult(
        {"health": {"overall_score": None}, "values": [float("nan"), np.float64("inf"), 0]}
    )
    payload = json.loads(result.to_json(), parse_constant=lambda value: pytest.fail(value))
    assert payload["values"] == [None, None, 0]
    assert payload["health"]["overall_score"] is None


def test_envelope_reports_coverage_without_certifying_missing_accuracy():
    e = build_envelope(
        {
            "tensor": {"matrix": {"approximate": True, "method": "sketch"}},
            "execution": {"sampled": True},
            "ml_readiness": {"score": None, "label": "not_applicable"},
        }
    )
    assert e["execution"]["coverage"] == "sampled"
    by_name = {d["name"]: d for d in e["diagnostics"]}
    assert by_name["tensor.matrix"]["accuracy"] == "approximate"
    assert by_name["ml_readiness"]["status"] == "not_applicable"
    assert by_name["tensor"]["accuracy"] == "unspecified"


def test_envelope_respects_planner_skips_and_failed_diagnostics():
    e = build_envelope(
        {
            "anomalies_v2": {"available": False},
            "execution": {
                "module_status": {"anomaly_detection": "disabled", "modeling": "not_applicable"}
            },
            "tensor": {"matrix": {"error": "failure"}},
        }
    )
    states = {d["name"]: d["status"] for d in e["diagnostics"]}
    assert states["anomalies_v2"] == "skipped"
    assert states["model_leaderboard"] == "not_applicable"
    assert states["tensor.matrix"] == "failed"
    assert e["status"] == "partial"


def test_tensor_dispatch_does_not_import_tabular_or_modeling_stack():
    code = "import sys,numpy as np; from framevitals import prism; prism(np.eye(3)); assert 'framevitals.pipeline' not in sys.modules; assert 'sklearn' not in sys.modules"
    subprocess.run([sys.executable, "-c", code], check=True, capture_output=True)


def test_cli_envelope_agrees_with_python(tmp_path):
    path = tmp_path / "values.npy"
    np.save(path, np.eye(3))
    completed = subprocess.run(
        [
            sys.executable,
            "-m",
            "framevitals",
            "prism",
            str(path),
            "--depth",
            "quick",
            "--format",
            "envelope",
        ],
        capture_output=True,
        text=True,
        check=True,
    )
    actual = json.loads(completed.stdout)
    expected = prism(path, depth="quick", artifacts=False).to_envelope()
    assert actual["analysis"] == expected["analysis"]


def test_invalid_depth_and_unsafe_numpy_fail_clearly(tmp_path):
    with pytest.raises(ValueError, match="depth"):
        prism(np.eye(3), depth="typo")
    path = tmp_path / "unsafe.npy"
    np.save(path, np.array([{"unsafe": True}], dtype=object))
    with pytest.raises(ValueError):
        prism(path)
    result = subprocess.run(
        [sys.executable, "-m", "framevitals", "prism", str(path)], capture_output=True, text=True
    )
    assert result.returncode == 2 and "Traceback" not in result.stderr


def test_archive_and_xml_preflight(tmp_path):
    path = tmp_path / "unsafe.docx"
    with ZipFile(path, "w") as z:
        z.writestr("../escaped.xml", "bad")
    with pytest.raises(ValueError, match="archive member"):
        validate_input(path)
    xml = tmp_path / "graph.graphml"
    xml.write_text('<!DOCTYPE x [<!ENTITY a "boom">]><x>&a;</x>')
    with pytest.raises(ValueError, match="unsafe XML"):
        validate_input(xml)


@pytest.fixture
def web(tmp_path, monkeypatch):
    pytest.importorskip("flask")
    from framevitals import webapp

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(webapp, "UPLOAD_DIR", Path("uploads"))
    with webapp.REPORT_LOCK:
        webapp.ANALYSIS_CACHE.clear()
        webapp.REPORT_JOBS.clear()
        webapp.UPLOAD_PATHS.clear()
    return webapp.app.test_client()


def upload(client, endpoint, content, filename, **options):
    return client.post(
        endpoint,
        data={"dataset": (io.BytesIO(content), filename), "analysis_mode": "quick", **options},
    )


@pytest.mark.parametrize(
    ("name", "content", "kind"),
    [
        ("data.csv", b"x,y\n1,2\n3,4\n", "tabular"),
        ("data.tsv", b"x\ty\n1\t2\n", "tabular"),
        ("data.json", b'{"a": [1,null,2]}', "nested"),
        ("data.jsonl", b'{"a":1}\n{"a":2}\n', "nested"),
        ("data.yaml", b"a: [1, 2]", "nested"),
        ("data.toml", b"a = [1, 2]", "nested"),
        ("data.md", b"# Heading\nSome document text.", "document"),
        ("data.html", b"<h1>Heading</h1><script>unsafe()</script><p>Safe text</p>", "document"),
    ],
)
def test_browser_source_workflow_and_cleanup(web, name, content, kind):
    preview = upload(web, "/api/inspect", content, name)
    assert preview.status_code == 200, preview.get_json()
    assert preview.json["source"]["kind"] == kind
    result = upload(web, "/api/prism", content, name)
    assert result.status_code == 200, result.get_json()
    assert result.json["source"]["kind"] == kind
    assert result.json["filename"] == name
    assert result.json["analysis"]["filename"] == name
    assert not list(Path("uploads").glob("*"))


def test_http_tensor_and_python_agree(web, tmp_path):
    path = tmp_path / "tensor.npy"
    np.save(path, np.eye(4))
    response = upload(web, "/api/prism", path.read_bytes(), path.name)
    assert response.status_code == 200
    expected = prism(path, depth="quick").analysis
    assert response.json["analysis"]["tensor"] == expected["tensor"]
    assert response.json["analysis"]["health"] == expected["health"]


def test_web_errors_and_ownership(web):
    response = upload(web, "/api/prism", b"unsafe", "unsafe.pkl")
    assert response.status_code == 400
    assert response.is_json
    response = upload(web, "/api/prism", b"{bad", "bad.json")
    assert response.status_code == 400
    response = upload(web, "/api/prism", b"a,b\n1,2\n", "data.csv")
    identifier = response.json["id"]
    from framevitals.webapp import app

    other = app.test_client()
    for url in [
        f"/download-cleaned/{identifier}",
        f"/download-report/{identifier}",
        f"/api/report-status/{identifier}",
    ]:
        assert other.get(url).status_code == 403
    assert web.post("/api/ai-report", json={"dataset_id": identifier}).status_code == 403
    assert (
        other.post("/api/ai-report", json={"dataset_id": identifier, "allow_ai": True}).status_code
        == 403
    )
    assert web.post("/api/ask", json=[]).status_code == 400


def test_cross_origin_post_and_oversized_request_are_json(web):
    assert (
        web.post("/api/prism", headers={"Origin": "https://untrusted.example"}).status_code == 403
    )
    from framevitals.webapp import app

    with pytest.MonkeyPatch.context() as m:
        m.setitem(app.config, "MAX_CONTENT_LENGTH", 10)
        response = upload(web, "/api/prism", b"x" * 100, "data.txt")
        assert response.status_code == 413 and response.is_json


def test_format_catalog_matches_public_registry(web):
    catalog = web.get("/api/capabilities").json
    assert catalog["formats"] == json.loads(json.dumps(capabilities()["formats"]))
    assert all(".pkl" not in spec["extensions"] for spec in catalog["formats"])


def test_browser_graph_and_bounded_preview(web, tmp_path):
    nx = pytest.importorskip("networkx")
    path = tmp_path / "network.graphml"
    nx.write_graphml(nx.path_graph(100), path)
    result = upload(web, "/api/prism", path.read_bytes(), path.name)
    assert result.status_code == 200, result.json
    preview = result.json["analysis"]["graph"]["topology_preview"]
    assert len(preview["nodes"]) == 48
    assert len(preview["edges"]) <= 120
    assert preview["sampled"] is True
    assert result.json["analysis"]["graph"]["nodes"] == 100


@pytest.mark.parametrize(
    "suffix", [".pdf", ".docx", ".pptx", ".xlsx", ".parquet", ".safetensors", ".onnx"]
)
def test_optional_binary_formats_through_http(web, tmp_path, suffix):
    path = tmp_path / f"source{suffix}"
    if suffix == ".pdf":
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(100, 100)
        writer.write(path)
    elif suffix == ".docx":
        from docx import Document

        document = Document()
        document.add_paragraph("Source text")
        document.save(path)
    elif suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Source text"
        presentation.save(path)
    elif suffix == ".xlsx":
        import pandas as pd

        pd.DataFrame({"a": [1, 2], "b": [3, 4]}).to_excel(path, index=False)
    elif suffix == ".parquet":
        pytest.importorskip("pyarrow")
        import pandas as pd

        pd.DataFrame({"a": [1, 2], "b": [3, 4]}).to_parquet(path)
    elif suffix == ".safetensors":
        import struct

        header = json.dumps(
            {"weight": {"dtype": "F32", "shape": [2], "data_offsets": [0, 8]}}
        ).encode()
        path.write_bytes(
            struct.pack("<Q", len(header)) + header + np.array([1.0, 2.0], dtype="<f4").tobytes()
        )
    else:
        onnx = pytest.importorskip("onnx")
        from onnx import TensorProto, helper

        graph = helper.make_graph(
            [helper.make_node("Identity", ["x"], ["y"])],
            "test",
            [helper.make_tensor_value_info("x", TensorProto.FLOAT, [1, 2])],
            [helper.make_tensor_value_info("y", TensorProto.FLOAT, [1, 2])],
        )
        onnx.save(helper.make_model(graph), path)
    result = upload(web, "/api/prism", path.read_bytes(), path.name)
    assert result.status_code == 200, result.json
    assert result.json["schema_version"] == "1"
    assert not list(Path("uploads").glob("*"))


def test_blank_pdf_is_not_reported_as_fully_understood(web, tmp_path):
    from pypdf import PdfWriter

    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(100, 100)
    writer.write(path)
    response = upload(web, "/api/prism", path.read_bytes(), path.name)
    info = response.json["analysis"]["document"]
    assert info["empty_text"] is True and info["ocr_performed"] is False
    assert response.json["findings"]


def test_capacity_limit_rejects_without_leaving_uploads(web):
    from framevitals.webapp import ANALYSIS_SLOTS

    ANALYSIS_SLOTS.acquire()
    ANALYSIS_SLOTS.acquire()
    try:
        response = upload(web, "/api/prism", b"hello", "notes.txt")
        assert response.status_code == 429
    finally:
        ANALYSIS_SLOTS.release()
        ANALYSIS_SLOTS.release()


def test_html_missing_score_is_unavailable():
    from framevitals.reporting.html import render_analysis_html
    result = AnalysisResult({'filename': 'missing.csv', 'health': {'overall_score': None}, 'ml_readiness': {'score': None}})
    html = render_analysis_html(result)
    assert 'Unavailable' in html
    assert '0/100' not in html
